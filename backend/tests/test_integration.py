import io
from datetime import date,timedelta
from docx import Document
from backend import store as db
from backend.seed import seed
from backend.services import perform_task,update_runs,process_refreshes
from backend.worker import claim

def drain():
    for _ in range(200):
        t=claim()
        if not t:break
        perform_task(t['id'])
    update_runs()

def test_full_demo_sources_evidence_and_scoring(client):
    seed();drain()
    result=client.get('/api/jobs/job_demo_1/matches').json()
    assert len(result['rows'])==26
    assert result['counts']['excluded']==4
    assert {a['profile']['source_type'] for a in result['rows']}=={'current_applicant','past_applicant','employee'}
    for a in result['rows']:
        assert 0<=a['overall_score']<=100
        ids={e['id'] for e in a['profile']['evidence']}
        assert all(e in ids for r in a['results'] for e in r['evidence_ids'])
    assert any(a['overlooked'] for a in result['rows'])

def test_interview_removes_cached_result_and_exports(client):
    seed();drain()
    rows=client.get('/api/jobs/job_demo_1/matches').json()['rows']
    a=next(a for a in rows if a['profile']['source_type']=='past_applicant')
    response=client.patch('/api/jobs/job_demo_1/profiles/'+a['profile_id']+'/review',json={'revision':0,'status':'shortlisted','reason':'Relevant experience reviewed','move_to_interview':True})
    assert response.status_code==200
    assert client.get('/api/assessments/'+a['id']).status_code==410
    assert a['profile_id'] not in [x['profile_id'] for x in client.get('/api/jobs/job_demo_1/matches').json()['rows']]
    assert a['profile']['facts']['name'] not in client.get('/api/jobs/job_demo_1/exports?format=csv').text

def test_weight_only_reassessment_reuses_all_interpretations(client):
    seed();drain()
    j=client.get('/api/jobs/job_demo_1').json();cs=j['draft'];cs[0]['weight']=3
    saved=client.put('/api/jobs/job_demo_1/rubric-draft',json={'revision':j['revision'],'criteria':cs}).json()
    approved=client.post('/api/jobs/job_demo_1/rubric-versions',json={'revision':saved['revision'],'overlap_acknowledged':True})
    assert approved.status_code==200
    drain()
    r=client.get('/api/processing/'+approved.json()['run']['id']).json()
    assert r['reused']==26
    assert client.get('/api/stats').json()['model_calls']==0

def test_stale_edits_and_same_rubric_comparison(client):
    seed();drain();j=client.get('/api/jobs/job_demo_1').json()
    bad=client.put('/api/jobs/job_demo_1/rubric-draft',json={'revision':0,'criteria':j['draft']})
    assert bad.status_code==409
    rows=client.get('/api/jobs/job_demo_1/matches').json()['rows']
    ids=[a['id'] for a in rows[:3]]
    assert client.post('/api/jobs/job_demo_1/comparisons',json={'assessment_ids':ids}).status_code==200
    scenario=client.post('/api/assessments/'+ids[0]+'/scenarios',json={'target':100})
    assert scenario.status_code==200

def test_delete_cascades_and_permissions(client,monkeypatch):
    seed();drain()
    import backend.main as main
    monkeypatch.setattr(main,'POOLS',['current_applicant'])
    assert client.delete('/api/profiles/profile_demo_11').status_code==403
    assert all(p['source_type']=='current_applicant' for p in client.get('/api/profiles').json())
    assert client.delete('/api/profiles/profile_demo_1').status_code==200
    with db.Session() as s:
        assert not db.get(s,'profile_demo_1')
        assert not any(a['profile_id']=='profile_demo_1' for a in db.all_of(s,'assessment'))
        assert not any(a.get('profile_id')=='profile_demo_1' for a in db.all_of(s,'profile_version'))

def test_queue_lease_recovery_and_per_candidate_failure(client):
    seed()
    t=claim()
    with db.Session.begin() as s:
        current=db.get(s,t['id']);db.put(s,t['id'],{**current,'lease_until':0})
    reclaimed=claim();assert reclaimed['id']==t['id']
    perform_task(t['id']);drain()
    assert client.get('/api/processing').json()['runs'][0]['status']=='completed'

def test_import_correction_and_refresh(client):
    j=client.post('/api/jobs',json={'title':'Java developer','description':'Develop Java services and PostgreSQL applications.','internal':True}).json()
    approval=client.post('/api/jobs/'+j['id']+'/rubric-versions',json={'revision':j['revision'],'overlap_acknowledged':True});assert approval.status_code==200
    d=Document();d.add_paragraph('Synthetic Test Candidate');d.add_paragraph('Developer | 2022-01 to present. Built Java services and PostgreSQL applications.')
    b=io.BytesIO();d.save(b)
    result=client.post('/api/imports',data={'job_id':j['id'],'source_type':'current_applicant'},files=[('files',('candidate.docx',b.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))],headers={'Idempotency-Key':'import-1'})
    assert result.status_code==200
    p=client.get('/api/profiles').json()[0]
    drain();process_refreshes();drain()
    assert len(client.get('/api/jobs/'+j['id']+'/matches').json()['rows'])==1
    p=client.get('/api/profiles').json()[0]
    corrected=client.patch('/api/profiles/'+p['id'],json={'revision':p['revision'],'reason':'Correct source text','correction_text':'Synthetic Test Candidate\nBuilt Java services.'})
    assert corrected.status_code==200
    assert corrected.json()['version_id']!=p['version_id']
    assert client.get('/api/jobs/'+j['id']+'/matches').json()['rows']==[]
    process_refreshes();drain()
    assert len(client.get('/api/jobs/'+j['id']+'/matches').json()['rows'])==1
